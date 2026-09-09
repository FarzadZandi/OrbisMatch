#!/usr/bin/env python3
"""Offline audit/rebuild pipeline for Orbis candidate matches."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import openpyxl

from matching_logic import score_candidate, select_best_match


DELIVERABLE_COLUMNS = [
    "Identifier",
    "Score",
    "Matched BvD ID",
    "Matched company name",
    "Matched via",
    "Comments",
]

REFERENCE_DELIVERABLE_COLUMNS = [
    "Order",
    "Own ID",
    "Company name",
    "City",
    "Country",
    "ISO2",
    "Website",
    "HQ city",
    "Zipcode",
    "Trade register number",
    "NationalID_norm",
    "Cand_BvD_guess",
    "Founders",
    "Trade register name",
    "Industries",
    "Sub industries",
    "Launch year",
]

AUDIT_COLUMNS = [
    "Own ID",
    "Reference company",
    "Reference country",
    "Reference city",
    "Reference website",
    "Candidate index",
    "Candidate BvD ID",
    "Candidate name",
    "Candidate city",
    "Candidate country",
    "Candidate national ID",
    "Orbis score letter",
    "Orbis score numeric",
    "Full address",
    "Domain",
    "Orbis ID",
    "Company status",
    "Owner name",
    "Management",
    "Was selected",
    "Matched via",
    "Witnesses matched",
    "Reason 1",
    "Reason 2",
    "Collection status",
    "Incomplete reason",
]


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def string_value(value: Any) -> str:
    return "" if value is None else str(value)


def resolve_input_path(path: Path) -> Path:
    if path.exists():
        return path
    fallback = Path(path.name)
    if fallback.exists():
        return fallback
    return path


def load_reference(path: Path) -> dict[str, dict[str, Any]]:
    path = resolve_input_path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_cell(value) for value in next(rows)]
    own_index = headers.index("Own ID")
    reference: dict[str, dict[str, Any]] = {}
    for raw in rows:
        own_id = normalize_cell(raw[own_index] if own_index < len(raw) else "")
        if not own_id:
            continue
        reference[own_id] = {
            header: raw[index] if index < len(raw) else ""
            for index, header in enumerate(headers)
            if header
        }
    workbook.close()
    return reference


def incomplete_count_from_decision(record: dict[str, Any]) -> int:
    if "incomplete_candidate_count" in record:
        try:
            return int(record.get("incomplete_candidate_count") or 0)
        except (TypeError, ValueError):
            return 0
    incomplete_candidates = record.get("incomplete_candidates")
    if isinstance(incomplete_candidates, list):
        return len(incomplete_candidates)
    return 0


def load_candidates(path: Path) -> tuple[dict[str, list[dict[str, Any]]], dict[str, dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    decisions_meta_by_id: dict[str, dict[str, Any]] = {}
    if not path.exists():
        return grouped, decisions_meta_by_id
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            record = json.loads(line)
            own_id = str(record.get("own_id") or "").strip()
            if not own_id:
                continue
            if record.get("source") == "decision":
                incomplete_candidates = record.get("incomplete_candidates")
                decisions_meta_by_id[own_id] = {
                    "incomplete_candidate_count": incomplete_count_from_decision(record),
                    "incomplete_candidates": incomplete_candidates if isinstance(incomplete_candidates, list) else [],
                    "decision_type": record.get("decision_type", ""),
                    "auto_match_candidate": record.get("auto_match_candidate") if isinstance(record.get("auto_match_candidate"), dict) else None,
                    "matched_via": record.get("matched_via", ""),
                    "reason_1": record.get("reason_1", ""),
                    "reason_2": record.get("reason_2", ""),
                    "comments": record.get("comments", ""),
                }
                continue
            grouped.setdefault(own_id, []).append(record)
    return grouped, decisions_meta_by_id


def build_decisions(
    reference: dict[str, dict[str, Any]],
    candidates_by_id: dict[str, list[dict[str, Any]]],
    decisions_meta_by_id: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    all_ids = set(reference) | set(candidates_by_id) | set(decisions_meta_by_id)
    return {
        own_id: build_auto_match_decision(decisions_meta_by_id[own_id])
        if decisions_meta_by_id.get(own_id, {}).get("decision_type") == "orbis_national_id_auto_match"
        else select_best_match(
            reference.get(own_id),
            candidates_by_id.get(own_id, []),
            incomplete_candidate_count=int(decisions_meta_by_id.get(own_id, {}).get("incomplete_candidate_count", 0)),
        )
        for own_id in all_ids
    }


def build_auto_match_decision(meta: dict[str, Any]) -> Any:
    return SimpleNamespace(
        selected=meta.get("auto_match_candidate") or {},
        evidence=None,
        decided=True,
        matched_via=meta.get("matched_via") or "Orbis National ID auto-match",
        reason_1=meta.get("reason_1") or "Matched automatically by Orbis Batch Search using the supplied National ID (trade register number).",
        reason_2=meta.get("reason_2") or "Remaining candidates were not collected and our two-witness rule was not independently applied.",
        comments=meta.get("comments") or (
            "Matched automatically by Orbis Batch Search using the supplied National ID (trade register number). "
            "Remaining candidates were not collected and our two-witness rule was not independently applied."
        ),
        best_candidate_if_no_match=None,
        tied_candidates=None,
    )


def header_map(sheet: Any) -> dict[str, int]:
    return {
        normalize_cell(cell.value): cell.column
        for cell in sheet[1]
        if normalize_cell(cell.value)
    }


def ensure_columns(sheet: Any, columns: list[str]) -> dict[str, int]:
    headers = header_map(sheet)
    next_col = sheet.max_column + 1
    for column in columns:
        if column not in headers:
            sheet.cell(row=1, column=next_col, value=column)
            headers[column] = next_col
            next_col += 1
    return headers


def decision_comments(decision: Any) -> str:
    return " ".join(
        part.strip()
        for part in [decision.reason_1, decision.reason_2, decision.comments]
        if part and part.strip()
    )


def tied_candidates_text(tied_candidates: list[dict[str, Any]]) -> str:
    parts = []
    for candidate in tied_candidates:
        name = string_value(candidate.get("candidate_name", ""))
        bvdid = string_value(candidate.get("bvdid", ""))
        if name and bvdid:
            parts.append(f"{name} ({bvdid})")
        elif name:
            parts.append(name)
        elif bvdid:
            parts.append(f"({bvdid})")
    return "; ".join(parts)


def deliverable_values(decision: Any) -> list[str]:
    if decision and decision.selected:
        selected = decision.selected
        values = [
            selected.get("national_id") or selected.get("identifier") or "",
            selected.get("score_letter") or selected.get("score_numeric") or "",
            selected.get("bvdid") or "",
            selected.get("candidate_name") or "",
            decision.matched_via,
            decision_comments(decision),
        ]
    elif decision and getattr(decision, "tied_candidates", None):
        tie_text = tied_candidates_text(decision.tied_candidates)
        comments = f"Undecided - manual review; tied candidates: {tie_text}. {decision_comments(decision)}".strip()
        values = [
            "",
            "",
            "Undecided",
            "Undecided",
            "Undecided (tie)",
            comments,
        ]
    else:
        values = [
            "",
            "",
            "No match",
            "No match",
            "No match",
            decision_comments(decision) if decision else "",
        ]
    return [string_value(value) for value in values]


def write_deliverable(
    reference: dict[str, dict[str, Any]],
    decisions_by_own_id: dict[str, Any],
    output_path: Path,
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Matched companies"
    sheet.append(REFERENCE_DELIVERABLE_COLUMNS + DELIVERABLE_COLUMNS)

    for own_id, ref_row in reference.items():
        reference_values = [string_value(ref_row.get(column, "")) for column in REFERENCE_DELIVERABLE_COLUMNS]
        sheet.append(reference_values + deliverable_values(decisions_by_own_id.get(own_id)))

    extra_ids = sorted(
        set(decisions_by_own_id) - set(reference),
        key=lambda value: int(value) if value.isdigit() else value,
    )
    own_id_index = REFERENCE_DELIVERABLE_COLUMNS.index("Own ID")
    for own_id in extra_ids:
        reference_values = ["" for _column in REFERENCE_DELIVERABLE_COLUMNS]
        reference_values[own_id_index] = own_id
        sheet.append(reference_values + deliverable_values(decisions_by_own_id.get(own_id)))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def selected_key(candidate: dict[str, Any] | None) -> tuple[str, str]:
    if not candidate:
        return ("", "")
    return (str(candidate.get("bvdid") or ""), str(candidate.get("candidate_index") or ""))


def candidate_key(candidate: dict[str, Any]) -> tuple[str, str]:
    return (str(candidate.get("bvdid") or ""), str(candidate.get("candidate_index") or ""))


def write_audit(
    candidates_by_id: dict[str, list[dict[str, Any]]],
    decisions_by_own_id: dict[str, Any],
    decisions_meta_by_id: dict[str, dict[str, Any]],
    reference: dict[str, dict[str, Any]],
    output_path: Path,
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Candidates audit"
    sheet.append(AUDIT_COLUMNS)

    all_ids = sorted(
        set(reference) | set(candidates_by_id) | set(decisions_meta_by_id),
        key=lambda value: int(value) if value.isdigit() else value,
    )
    for own_id in all_ids:
        ref = reference.get(own_id, {})
        candidates = candidates_by_id.get(own_id, [])
        decision = decisions_by_own_id.get(own_id)
        decision_meta = decisions_meta_by_id.get(own_id, {})
        incomplete_candidates = decision_meta.get("incomplete_candidates", [])
        auto_match_candidate = decision_meta.get("auto_match_candidate")
        winner_key = selected_key(decision.selected) if decision and decision.selected else ("", "")

        if not candidates and auto_match_candidate:
            candidate = auto_match_candidate
            sheet.append(
                [
                    own_id,
                    string_value(ref.get("Company name", "")),
                    string_value(ref.get("ISO2") or ref.get("Country") or ""),
                    string_value(ref.get("City") or ref.get("HQ city") or ""),
                    string_value(ref.get("Website", "")),
                    string_value(candidate.get("candidate_index", "")),
                    string_value(candidate.get("bvdid", "")),
                    string_value(candidate.get("candidate_name", "")),
                    string_value(candidate.get("city", "")),
                    string_value(candidate.get("country", "")),
                    string_value(candidate.get("national_id") or candidate.get("identifier") or ""),
                    string_value(candidate.get("score_letter", "")),
                    string_value(candidate.get("score_numeric", "")),
                    string_value(candidate.get("full_address", "")),
                    string_value(candidate.get("domain") or candidate.get("website") or ""),
                    string_value(candidate.get("orbis_id", "")),
                    string_value(candidate.get("company_status", "")),
                    string_value(candidate.get("owner_name", "")),
                    string_value(candidate.get("management", "")),
                    "Yes",
                    string_value(decision.matched_via if decision else "Orbis National ID auto-match"),
                    "",
                    string_value(decision.reason_1 if decision else ""),
                    string_value(decision.reason_2 if decision else ""),
                    "Orbis auto-match (remaining candidates not collected)",
                    "",
                ]
            )
        elif not candidates:
            sheet.append(
                [
                    own_id,
                    string_value(ref.get("Company name", "")),
                    string_value(ref.get("ISO2") or ref.get("Country") or ""),
                    string_value(ref.get("City") or ref.get("HQ city") or ""),
                    string_value(ref.get("Website", "")),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "No",
                    "",
                    "",
                    "",
                    "",
                    "No collected candidates",
                    "",
                ]
            )

        for candidate in candidates:
            evidence = score_candidate(ref, candidate) if ref else None
            is_selected = bool(winner_key != ("", "") and candidate_key(candidate) == winner_key)
            sheet.append(
                [
                    own_id,
                    string_value(ref.get("Company name", "")),
                    string_value(ref.get("ISO2") or ref.get("Country") or ""),
                    string_value(ref.get("City") or ref.get("HQ city") or ""),
                    string_value(ref.get("Website", "")),
                    string_value(candidate.get("candidate_index", "")),
                    string_value(candidate.get("bvdid", "")),
                    string_value(candidate.get("candidate_name", "")),
                    string_value(candidate.get("city", "")),
                    string_value(candidate.get("country", "")),
                    string_value(candidate.get("national_id") or candidate.get("identifier") or ""),
                    string_value(candidate.get("score_letter", "")),
                    string_value(candidate.get("score_numeric", "")),
                    string_value(candidate.get("full_address", "")),
                    string_value(candidate.get("domain") or candidate.get("website") or ""),
                    string_value(candidate.get("orbis_id", "")),
                    string_value(candidate.get("company_status", "")),
                    string_value(candidate.get("owner_name", "")),
                    string_value(candidate.get("management", "")),
                    "Yes" if is_selected else "No",
                    string_value(decision.matched_via if is_selected and decision else ""),
                    "; ".join(evidence.witnesses) if evidence else "",
                    string_value(decision.reason_1 if is_selected and decision else ""),
                    string_value(decision.reason_2 if is_selected and decision else ""),
                    "Collected",
                    "",
                ]
            )

        for incomplete in incomplete_candidates:
            sheet.append(
                [
                    own_id,
                    string_value(ref.get("Company name", "")),
                    string_value(ref.get("ISO2") or ref.get("Country") or ""),
                    string_value(ref.get("City") or ref.get("HQ city") or ""),
                    string_value(ref.get("Website", "")),
                    string_value(incomplete.get("candidate_index", "")),
                    string_value(incomplete.get("bvdid", "")),
                    string_value(incomplete.get("candidate_name", "")),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    string_value(incomplete.get("company_status", "")),
                    string_value(incomplete.get("owner_name", "")),
                    "",
                    "No",
                    "",
                    "",
                    "",
                    "",
                    string_value(incomplete.get("status") or "incomplete"),
                    string_value(incomplete.get("error", "")),
                ]
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def run(reference_path: Path, results_path: Path, deliverable_path: Path, audit_path: Path) -> None:
    reference = load_reference(reference_path)
    candidates_by_id, decisions_meta_by_id = load_candidates(results_path)
    decisions_by_own_id = build_decisions(reference, candidates_by_id, decisions_meta_by_id)
    write_deliverable(reference, decisions_by_own_id, deliverable_path)
    write_audit(candidates_by_id, decisions_by_own_id, decisions_meta_by_id, reference, audit_path)
    print(f"Wrote deliverable: {deliverable_path}")
    print(f"Wrote audit: {audit_path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Rebuild final match Excel files from collected Orbis JSONL.")
    parser.add_argument("--reference", default="files/master_enriched_reference.xlsx")
    parser.add_argument("--results", default="orbis_run/results.jsonl")
    parser.add_argument("--deliverable", default="orbis_run/Matched_companies_Farzad.xlsx")
    parser.add_argument("--audit", default="orbis_run/candidates_audit.xlsx")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    run(
        Path(args.reference),
        Path(args.results),
        Path(args.deliverable),
        Path(args.audit),
    )
